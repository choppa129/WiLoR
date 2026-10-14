"""Build a source-backed, editable research update deck."""
from pathlib import Path
import json
import pandas as pd
import numpy as np
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
BASE = ROOT/'evaluation/scale_diagnostic'
ANNOTATED = ROOT/'evaluation/freihand_annotated_interventions_full'
OUT = ROOT/'reports/wilor_experiment_review_corrected'
TRACKED_OUT = ROOT/'reports/wilor_experiment_review'
OUT.mkdir(parents=True,exist_ok=True)
TRACKED_OUT.mkdir(parents=True,exist_ok=True)
prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
BG='F4F6FA'; INK='17243A'; TEAL='007F86'; MUTED='58677B'
sources=[]

def text(s,x,y,w,h,content,size=20,color=INK,bold=False):
    box=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
    tf=box.text_frame;tf.word_wrap=True
    for i,line in enumerate(content.split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.text=line;p.font.name='Aptos';p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=RGBColor.from_string(color)
        p.space_after=Pt(12)
    return box

def slide(title,subtitle,source):
    s=prs.slides.add_slide(prs.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=RGBColor.from_string(BG)
    text(s,.5,.22,12.3,.6,title,28,bold=True)
    text(s,.5,.95,12.3,.55,subtitle,15,MUTED)
    text(s,.5,7.08,11.9,.25,source,9,MUTED)
    text(s,12.35,7.04,.5,.3,str(len(prs.slides)),11,MUTED)
    sources.append(source)
    s.notes_slide.notes_text_frame.text='Evidence: '+source+'\nLocal project root: '+str(ROOT)
    return s

def table(s,headers,rows,y=1.8,h=3.4):
    shape=s.shapes.add_table(len(rows)+1,len(headers),Inches(.55),Inches(y),Inches(12.2),Inches(h))
    for i,row in enumerate([headers]+rows):
        for j,v in enumerate(row):
            c=shape.table.cell(i,j);c.text=str(v);c.margin_left=Inches(.12);c.margin_top=Inches(.12)
            c.fill.solid();c.fill.fore_color.rgb=RGBColor.from_string(INK if i==0 else ('FFFFFF' if i%2 else 'E8EEF4'))
            for p in c.text_frame.paragraphs:
                p.font.name='Aptos';p.font.size=Pt(16);p.font.bold=i==0;p.font.color.rgb=RGBColor.from_string('FFFFFF' if i==0 else INK)

def picture(s,path,y=1.75,maxh=3.7):
    with Image.open(path) as im: w,h=im.size
    scale=min(12.2/w,maxh/h)
    s.shapes.add_picture(str(path),Inches((13.333-w*scale)/2),Inches(y),width=Inches(w*scale),height=Inches(h*scale))

metrics=pd.read_csv(ROOT/'evaluation/freihand_full_comparison_corrected/metrics.csv').set_index('model')
pred=pd.concat([pd.read_csv(BASE/'predictions'/m/'metrics.csv') for m in ['official','trained','hamer']])
cam=pd.read_csv(BASE/'camera_swap/per_sample.csv')
bet=pd.read_csv(BASE/'betas_ablation/per_sample.csv')
pose=pd.read_csv(BASE/'pose_ablation/per_sample.csv')
annotated=pd.read_csv(ANNOTATED/'summary.csv').set_index('variant')
paired=pd.read_csv(ANNOTATED/'paired_summary.csv')
audit=pd.read_csv(ROOT/'evaluation/freihand_loss_gradient_audit/summary.csv').set_index('loss')

s=slide('WiLoR reconstruction: training and scale diagnosis','Research update | Official WiLoR, official HaMeR and a FreiHAND-trained WiLoR','Sources: local evaluation reports and saved model configurations')
text(s,.7,1.9,11.8,1.2,'What explains the smaller projected hands?',34,TEAL,True)
text(s,.7,3.35,11.8,2.6,'Training learns useful reconstruction, but official WiLoR remains stronger on the local validation split.\nFull annotation interventions identify pose—especially finger articulation—as the main in-domain gap.\nThe apparent demo shrinkage is not a systematic FreiHAND scale bias; camera and geometry are co-adapted.',23)
text(s,.7,6.3,11.8,.5,'Evidence: 10,240 annotated validation samples + 45 fixed hand proposals from 8 demo images.',18,MUTED)

s=slide('Three models, two different comparison settings','Distinguish full detection pipelines from controlled reconstruction experiments.','Sources: demo.py; hamer_official/code/demo.py; wilor/models/wilor.py')
table(s,['Component','Official HaMeR','Official WiLoR','Self-trained WiLoR'],[
['Demo localization','Person detector → ViTPose','YOLO hand detector','Same WiLoR detector'],
['Reconstruction','ViT + MANO head','Initial prediction + refinement','WiLoR architecture + refinement'],
['Weights used','Official released checkpoint','Official released checkpoint','ViTPose initialization; FreiHAND training'],
['Current ablations','Reference betas / final output','Reference before / after','Model under diagnosis']],h=3.8)
text(s,.65,6.0,12,.75,'For diagnosis, all three consume the same saved crop tensors. HaMeR here is a reconstruction reference, not its full detection pipeline.',18,TEAL)

s=slide('Self-training setup','Configuration facts, not a claim of reproducing the official training recipe.','Sources: logs/train/runs/wilor_freihand_full_3gpu/{model_config.yaml,dataset_config.yaml,exec_time.log}')
table(s,['Item','Configuration'],[
['Data','FreiHAND only; weight 1.0; train shards 000–119, validation 120–130'],
['Initialization','ViTPose backbone initialization; WiLoR reconstruction/refinement trained'],
['Compute','3 × B200; batch 8/GPU (24 total); FP16 mixed precision'],
['Optimization','AdamW; LR 1e-5; weight decay 1e-4'],
['Run settings','Configured max_steps=100,000; recorded execution time 2.64 h'],
['Betas loss','Configured weight 0.0005; magnitude alone cannot establish weak supervision']],h=4.35)
text(s,.65,6.3,12,.55,'Manual optimization uses generator + discriminator updates; max_steps is not independently verified as 100k generator updates.',14,MUTED)

s=slide('Local FreiHAND validation: learning works, a gap remains','10,240 historical samples. Reprojection corrected arithmetically; no new evaluation.','Source: evaluation/freihand_full_comparison_corrected/metrics.csv; projection_audit/report.md')
rows=[]
for k,label in [('initial','ViTPose init + random head'),('trained','Self-trained WiLoR'),('official','Official WiLoR')]:
    rows.append([label]+[f'{metrics.loc[k,c]:.3f}' for c in ['val/mpjpe_mm','val/pa_mpjpe_mm','val/pve_mm','val/reprojection_px']])
table(s,['Model','MPJPE mm','PA-MPJPE mm','PVE mm','Reproj.¹'],rows,h=2.6)
improve=100*(1-metrics.loc['trained','val/mpjpe_mm']/metrics.loc['initial','val/mpjpe_mm'])
text(s,.65,4.8,12,.8,f'MPJPE improves {improve:.1f}% over initialization; self-trained error is 2.28× the official checkpoint.',23,TEAL,True)
text(s,.65,5.8,12,1.1,'Internal shard split, not the official benchmark. Capture independence and official training overlap remain unaudited; no HaMeR evaluation in this table.\n¹ Corrected crop-pixel error: old values ×2. Training loss and 3D metrics are unchanged.',14,MUTED)

s=slide('Controlled diagnostic: identical reconstruction inputs','Freeze the input before interpreting differences in the output.','Sources: evaluation/scale_diagnostic/fixed_inputs/manifest.json; predictions/*/provenance.json')
text(s,.7,1.85,11.8,3.6,'45 detector proposals from 8 images; left/right label and crop transform saved.\nWiLoR and HaMeR crop tensors match exactly: maximum absolute difference = 0.\nFP32 inference, TF32 disabled, layer skipping disabled; input hashes verified.\n225 outputs: WiLoR official 90 + self-trained 90 + HaMeR 45.',23)
text(s,.7,5.5,11.8,1.1,'“Size” = diagonal of the projected joint bounding box. This depends on camera, shape and pose; it is not anatomical hand size. Proposals may include false positives.',19,TEAL)

s=slide('Refinement expands the self-trained projection','The discrepancy is already present before refinement.','Source: evaluation/scale_diagnostic/report_stage2.md; predictions/*/metrics.csv')
rows=[]
for model,stage in [('official','before'),('official','after'),('trained','before'),('trained','after'),('hamer','final')]:
    r=pred[(pred.model==model)&(pred.stage==stage)]
    rows.append([model,stage,f'{r.joint_extent_px.mean():.2f}',f'{r.mesh_diag.mean():.4f}',f'{r.beta_norm.mean():.3f}'])
table(s,['Model','Stage','Mean 2D extent px','Mean 3D diagonal m','Mean beta norm'],rows,h=3.45)
text(s,.65,5.6,12,1.1,'Trained before / official before: median 72.20% (45/45 smaller).\nTrained after / before: median 135.50%; after / official after: 93.61% (30/45 smaller).',20,TEAL)

s=slide('Demo camera swap raised a hypothesis','Keep geometry fixed and replace only the saved camera translation.','Source: evaluation/scale_diagnostic/camera_swap/{per_sample.csv,report.md}')
rows=[]
for g,c in [('official_after','official_after'),('trained_after','trained_after'),('trained_after','official_after'),('official_after','trained_after')]:
    r=cam[(cam.comparison=='final')&(cam.geometry==g)&(cam.camera==c)]
    rows.append([g.split('_')[0],c.split('_')[0],f'{r.joint_extent_ratio.median()*100:.2f}%'])
table(s,['Geometry (pose + shape)','Camera','Median extent / official'],rows,h=3.1)
text(s,.65,5.35,12,1.3,'The trained geometry reaches 98.83% of the official extent with the official camera, versus 93.61% with its own.\nThis visual-reference result is not a ground-truth accuracy test; the full annotated intervention later shows that the swap worsens reprojection.',20,TEAL)

s=slide('Camera swap: a fixed-sample example','test1_000 | Gray: official reference joints. Green: tested combination.','Source: evaluation/scale_diagnostic/camera_swap/test1_000.jpg')
picture(s,BASE/'camera_swap/test1_000.jpg')
text(s,.65,5.7,12,1.0,'Left to right: official / official; trained / trained; trained / official; official / trained.\nThis example illustrates the intervention; aggregate statistics use all 45 proposals.',18,MUTED)

s=slide('Betas swap: limited median effect on projected extent','Hold the trained pose and camera fixed; change only the 10 MANO shape coefficients.','Source: evaluation/scale_diagnostic/betas_ablation/{per_sample.csv,report.md}')
baseline=bet[bet.betas_source=='trained'].set_index('sample');rows=[]
for source in ['official','hamer','zero']:
    r=bet[bet.betas_source==source].set_index('sample')
    rows.append([source,f'{((r.joint_extent_px/baseline.joint_extent_px).median()-1)*100:+.2f}%',f'{((r.palm_width_m/baseline.palm_width_m).median()-1)*100:+.2f}%'])
table(s,['Betas source','Median projected extent change','Median palm-width change'],rows,h=2.6)
text(s,.65,4.7,12,1.7,'Official betas: only 15/45 samples move closer to official projected size.\nEffects vary: P10–P90 extent change ≈ −4.7% to +10.1%.\nA larger beta norm does not prove shape failure or justify increasing the beta loss weight.',21,TEAL)

s=slide('Demo pose swap raised an orientation lead','Hold the trained betas and camera fixed; replace global and/or local rotations.','Source: evaluation/scale_diagnostic/pose_ablation/{per_sample.csv,report.md}')
rows=[]
for v,label in [('official_orientation','Global orientation only'),('official_fingers','Finger articulation only'),('official_both','Both')]:
    r=pose[pose.variant==v].set_index('sample');b=pose[pose.variant=='baseline'].set_index('sample')
    closer=(abs(r.joint_ratio_to_official-1)<abs(b.joint_ratio_to_official-1)-1e-6).sum()
    rows.append([label,f'{(r.joint_ratio_to_trained.median()-1)*100:+.2f}%',f'{r.joint_ratio_to_official.median()*100:.2f}%',f'{closer}/45'])
table(s,['Replacement','Change vs trained','Extent / official','Closer to official'],rows,h=2.7)
text(s,.65,4.9,12,1.6,'Orientation-only replacement brings the paired median extent ratio to 98.76% of official.\nThis measures agreement with another model, not ground truth. Full annotated results show finger pose is the clearer accuracy limitation.\nChanges interact; medians and their percentages must not be added.',19,TEAL)

s=slide('Pose and betas: visual checks','Same sample and fixed display window; gray is the trained baseline, green is the intervention.','Sources: pose_ablation/test1_000.jpg; betas_ablation/test1_000.jpg (under evaluation/scale_diagnostic)')
text(s,.65,1.5,12,.3,'POSE: baseline | official orientation | official fingers | both',15,TEAL,True)
picture(s,BASE/'pose_ablation/test1_000.jpg',1.92,1.85)
text(s,.65,4.05,12,.3,'BETAS: trained | official WiLoR | official HaMeR | zero',15,TEAL,True)
picture(s,BASE/'betas_ablation/test1_000.jpg',4.45,1.85)
text(s,.65,6.45,12,.35,'Baseline reconstruction check: <9e-8 m in 3D; pose experiment reprojection error <0.000203 px.',14,MUTED)

s=slide('Full annotated intervention: absolute metrics','All 10,240 deterministic local validation samples; the 512-sample pilot conclusion holds.','Sources: evaluation/freihand_annotated_interventions_full/{summary.csv,provenance.json}')
rows=[]
for variant,label in [
    ('trained','Self-trained'),('official','Official WiLoR'),
    ('trained_with_gt_orientation','Trained + GT orientation'),
    ('trained_with_gt_fingers','Trained + GT fingers'),
    ('trained_with_gt_pose','Trained + GT full pose')]:
    r=annotated.loc[variant]
    rows.append([label,f"{r.reprojection_px_mean:.3f}",f"{r.mpjpe_mm_mean:.3f}",
                 f"{r.pa_mpjpe_mm_mean:.3f}",f"{r.pve_mm_mean:.3f}",
                 f"{100*r.extent_ratio_median:.2f}%"])
table(s,['Variant','Reproj px','MPJPE mm','PA-MPJPE','PVE mm','Extent / GT'],rows,h=3.5)
text(s,.65,5.75,12,1.0,'Self-trained extent is 98.34% of GT: no strong systematic in-domain shrinkage.\nOfficial WiLoR is still substantially more accurate across all four error metrics.',19,TEAL)

s=slide('Full annotated intervention: paired evidence','Interventions are diagnostic counterfactuals; lower errors are better.','Sources: evaluation/freihand_annotated_interventions_full/{paired_summary.csv,report.md}')
def paired_row(candidate,metric):
    return paired[(paired.candidate==candidate)&(paired.metric==metric)].iloc[0]
rows=[]
for candidate,metric,label in [
    ('official','reprojection_px','Official vs trained: reprojection'),
    ('trained_geometry_official_camera','reprojection_px','Official camera on trained geometry'),
    ('trained_with_gt_orientation','mpjpe_mm','GT orientation: MPJPE'),
    ('trained_with_gt_fingers','pa_mpjpe_mm','GT fingers: PA-MPJPE'),
    ('trained_with_gt_pose','mpjpe_mm','GT full pose: MPJPE'),
    ('trained_with_gt_pose','pve_mm','GT full pose: PVE')]:
    r=paired_row(candidate,metric)
    rows.append([label,f"{int(r.improved)}/10,240",f"{r.median_relative_change_pct:+.2f}%"])
table(s,['Paired comparison','Improved samples','Median relative change'],rows,h=3.55)
text(s,.65,5.8,12,.9,'Camera swap worsens reprojection by 5.65% median: geometry and camera are co-adapted.\nGT fingers cut PA-MPJPE by 46.75%; complete GT pose cuts MPJPE 52.57% and PVE 68.65%.',19,TEAL)

s=slide('Loss audit: explicit pose supervision is weak','64 fixed samples, 8 local batches of 8; FP32 diagnostic gradients, no optimizer update.','Sources: evaluation/freihand_loss_gradient_audit/{summary.csv,report.md,provenance.json}')
rows=[]
for key,label in [('keypoints_2d','2D keypoints'),('keypoints_3d','3D keypoints'),
                  ('global_orient','Global orientation'),('hand_pose','Finger pose'),
                  ('betas','Betas'),('adversarial','Adversarial')]:
    r=audit.loc[key]
    rows.append([label,f"{r.weight:.4g}",f"{100*r.weighted_fraction_mean:.2f}%",
                 f"{r.grad_all_generator_median:.4f}",f"{r.grad_refine_pose_head_median:.4f}",
                 f"{r.grad_refine_shape_head_median:.4f}",f"{r.grad_refine_camera_head_median:.4f}"])
table(s,['Loss','Weight','Loss share','All grad','Pose head','Shape head','Camera head'],rows,h=3.55)
text(s,.65,5.75,12,1.0,'2D + 3D keypoints supply 84.44% of the weighted objective. Global-orientation supervision is only 0.16%.\nBetas are active (6.64%; grad 0.231), so increasing beta weight is not the evidence-based first repair.',18,TEAL)

s=slide('Recommendation: add data before changing weights','Keep the first multi-dataset experiment interpretable by changing one factor at a time.','Sources: full annotated intervention and loss-gradient audit; proposed experiment, not yet run')
table(s,['Evidence','Implication'],[
['FreiHAND median extent / GT = 98.34%.','No strong systematic in-domain scale bias.'],
['The larger size discrepancy appears on demo images.','Domain coverage is the stronger first hypothesis.'],
['All configured losses are active; betas contribute 6.64%.','No failed or clearly absent loss requires an immediate weight fix.'],
['Official WiLoR used broader training data.','Add WHIM alongside FreiHAND before tuning the objective.']],h=3.0)
text(s,.65,4.95,12,1.55,'Protocol: same ViTPose initialization, architecture, optimizer and current loss weights; change only the data mixture.\nStart with a 1,000-step FreiHAND + WHIM smoke run, verify sampling and validation, then scale up only if metrics and fixed-demo outputs improve.\nJudge MPJPE, PA-MPJPE, PVE and reprojection together—not projected extent alone.',18,TEAL)

s=slide('Conclusions and the next decision','Current evidence supports a controlled multi-dataset experiment before loss tuning.','Sources: full annotated intervention, loss-gradient audit and all local scale_diagnostic reports')
table(s,['Supported by current experiments','Still unresolved'],[
['Refinement usually compensates for smaller demo projections.','Why the demo domain shows a larger size discrepancy.'],
['No strong systematic FreiHAND scale shrinkage (98.34% median extent / GT).','Whether more varied training data fixes out-of-domain scale behavior.'],
['Pose—especially fingers—is the main annotated accuracy gap.','The best FreiHAND/WHIM sampling ratio and coverage.'],
['Camera and geometry are co-adapted; isolated swapping is misleading.','Per-sample failures, false detections and domain shift.']],h=3.15)
text(s,.65,5.25,12,1.35,'Next: add WHIM and run a 1,000-step FreiHAND + WHIM smoke test with the current loss weights unchanged.\nUse the same ViTPose initialization, architecture and optimizer so the data effect remains identifiable.\nSelect by full validation MPJPE, PA-MPJPE, PVE and reprojection—not projected extent alone.',19,TEAL)

prs.save(OUT/'WiLoR_experiments_and_scale_diagnosis.pptx')
prs.save(TRACKED_OUT/'WiLoR_experiments_and_scale_diagnosis.pptx')
check=Presentation(OUT/'WiLoR_experiments_and_scale_diagnosis.pptx')
assert len(check.slides)==16
for s in check.slides:
    for sh in s.shapes:
        assert sh.left>=0 and sh.top>=0
        assert sh.left+sh.width<=check.slide_width+100 and sh.top+sh.height<=check.slide_height+100
(OUT/'sources.json').write_text(json.dumps(sources,indent=2))
print('Created and structurally verified:',OUT/'WiLoR_experiments_and_scale_diagnosis.pptx')
